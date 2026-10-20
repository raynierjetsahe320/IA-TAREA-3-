import csv
import html as html_lib
import hashlib
import json
import math
import re
import random
import statistics
import time
import textwrap
from collections.abc import Callable
from collections import defaultdict
from pathlib import Path
from typing import Any

from .models import GAConfig, SAConfig, SolverResult, result_parameters
from .problems import (
    GraphInstance,
    KnapsackInstance,
    eggholder,
    knapsack_totals,
    queen_conflicts,
    value_density,
)
from .solvers import (
    coloring_ga_min,
    coloring_hybrid_min,
    coloring_sa_min,
    eggholder_ga,
    eggholder_sa,
    hybrid_ga_sa,
    knapsack_ga,
    knapsack_sa,
    queens_ga,
    queens_sa,
    solve_coloring_min_backtracking,
    solve_knapsack_backtracking,
    solve_queens_backtracking,
)

RESULT_FIELDS = [
    "problem",
    "instance",
    "instance_source",
    "instance_data",
    "instance_size",
    "algorithm",
    "run",
    "seed",
    "objective",
    "quality",
    "evaluations",
    "iterations",
    "time_seconds",
    "feasible",
    "limit_reached",
    "conflicts",
    "total_weight",
    "total_value",
    "solution",
    "parameters",
    "feasibility_status",
    "quality_gap_percent",
    "quality_status",
]

PROBLEM_LABELS = {
    "queens": "N-Reinas",
    "knapsack": "Mochila",
    "coloring": "Coloreado de grafos",
    "eggholder": "Eggholder",
}
ALGORITHM_LABELS = {
    "Backtracking": "Búsqueda exacta (backtracking)",
    "Simulated Annealing": "Recocido simulado",
    "Genetic Algorithm": "Algoritmo genético",
    "Hybrid GA+SA": "Híbrido genético + recocido",
}
PROBLEM_EXPLANATIONS = [
    (
        "N-Reinas",
        "Cada tablero se representa como una permutación: la posición de la lista indica la fila y el valor, la columna. Así queda garantizada una reina por fila y por columna; solo hay que minimizar ataques diagonales. En SA, intercambiar dos columnas conserva esa validez estructural; en GA, el cruce OX y la mutación por intercambio conservan la permutación. Backtracking elige primero la fila con menos columnas disponibles y prueba antes la columna que restringe menos a las demás.",
        "Esto reduce conflictos innecesarios y hace que una solución con cero ataques sea directamente válida. Al crecer N, las permutaciones posibles aumentan rápidamente; con un presupuesto corto, las metaheurísticas pueden quedarse con conflictos diagonales aunque su representación siempre evite columnas repetidas.",
    ),
    (
        "Mochila 0/1",
        "Un vector binario representa si se toma cada objeto (1) o se deja (0). Es la representación natural porque cada artículo es indivisible. SA cambia un bit; GA mezcla decisiones objeto por objeto. Ambos reparan el vector quitando artículos de menor valor por peso hasta respetar la capacidad. Backtracking ordena por valor/peso y usa una cota optimista fraccional para descartar ramas que ya no pueden superar la mejor selección.",
        "La reparación garantiza factibilidad, pero puede quitar un objeto que forme parte de la mejor combinación: los artículos no se pueden fraccionar y una decisión afecta al peso restante. Por eso factible no significa óptimo. En las instancias incluidas hay pocos objetos, de modo que backtracking suele poder comparar combinaciones y certificar el mejor valor si no alcanza el límite.",
    ),
    (
        "Coloreado de grafos",
        "Cada posición del vector representa un vértice y su entero es el color asignado. La calidad se mide primero por aristas en conflicto; una asignación solo es factible con cero conflictos. Para minimizar colores se parte de una coloración voraz DSATUR y se prueban cantidades menores, usando como límite inferior una clique encontrada. SA recolorea un vértice; GA combina colores por vértice; el híbrido aplica GA y luego refina con SA.",
        "Una asignación completa puede ser infactible durante la búsqueda; por eso obtener pocos colores no basta si quedan vecinos del mismo color. Los grafos grandes tienen muchas interacciones entre vértices y demostrar que no existe una coloración con un color menos puede requerir una búsqueda combinatoria muy grande. Una solución válida de cuatro colores es un buen resultado observado, pero no es un certificado de que cuatro sea el mínimo si backtracking quedó truncado.",
    ),
    (
        "Eggholder",
        "La solución es un par de coordenadas reales dentro de [-512, 512] en cada dimensión. SA hace desplazamientos gaussianos locales con tamaño ligado al enfriamiento; GA usa cruza lineal y muta una coordenada con ruido. Los límites se recortan para que las propuestas sigan dentro del dominio.",
        "La función tiene muchos valles y picos. Un paso local pequeño puede atascarse en un mínimo local; pasos grandes exploran más pero pueden saltarse una zona prometedora. GA mantiene varias propuestas, aunque su población y generaciones limitadas no garantizan hallar el mínimo global. Aquí estar dentro de los límites es factibilidad; la calidad depende de qué tan bajo sea el valor obtenido.",
    ),
]


def _goal_label(problem: str) -> str:
    return {
        "queens": "menos conflictos",
        "knapsack": "más valor en la mochila",
        "coloring": "menos colores, sin conflictos",
        "eggholder": "menor valor de la función",
    }.get(problem, "mejor resultado")


def _format_quality(problem: str, value: Any) -> str:
    if value == "" or value is None:
        return "sin solución factible"
    number = float(value)
    formatted = f"{number:.3f}".rstrip("0").rstrip(".")
    suffix = {
        "queens": " conflictos",
        "knapsack": " de valor",
        "coloring": " colores",
        "eggholder": " puntos",
    }.get(problem, "")
    return f"{formatted}{suffix}"


def _format_seconds(value: float) -> str:
    return "<0.001 s" if value < 0.001 else f"{value:.3f} s"


def _quality_gap(problem: str, value: float, reference: float) -> float:
    if math.isclose(value, reference, rel_tol=1e-12, abs_tol=1e-12):
        return 0.0
    scale = abs(reference) or abs(value) or 1.0
    difference = reference - value if problem == "knapsack" else value - reference
    return max(0.0, difference / scale * 100)


def _quality_status(problem: str, gap: float | None, aggregate: bool = False) -> str:
    if gap is None:
        return "Sin solución factible; calidad no comparable"
    if gap <= 1e-9:
        if problem == "queens":
            return "Factible y de calidad óptima observada: 0 conflictos"
        return (
            "Empata el mejor promedio factible observado"
            if aggregate
            else "Empata la mejor calidad factible individual observada"
        )
    reference = (
        "mejor promedio factible observado"
        if aggregate
        else "mejor calidad factible individual observada"
    )
    return f"Desventaja de {gap:.1f}% frente al {reference}"


def _annotate_record_quality(
    records: list[dict[str, Any]],
) -> None:
    references: dict[tuple[str, str], float] = {}
    for record in records:
        if not record["feasible"] or record["quality"] is None:
            continue
        key = (record["problem"], record["instance"])
        value = float(record["quality"])
        current = references.get(key)
        if current is None:
            references[key] = value
        elif record["problem"] == "knapsack":
            references[key] = max(current, value)
        else:
            references[key] = min(current, value)

    for record in records:
        record["feasibility_status"] = (
            "Factible: respeta las restricciones"
            if record["feasible"]
            else "No factible: incumple al menos una restricción"
        )
        reference = references.get((record["problem"], record["instance"]))
        quality = record["quality"]
        gap = (
            _quality_gap(record["problem"], float(quality), reference)
            if record["feasible"] and quality is not None and reference is not None
            else None
        )
        record["quality_gap_percent"] = "" if gap is None else round(gap, 6)
        record["quality_status"] = _quality_status(record["problem"], gap)


def _parameter_explanations(parameters: dict[str, Any]) -> list[str]:
    sa = parameters.get("sa", {})
    ga = parameters.get("ga", {})
    sa_iterations = int(sa.get("iterations", 0))
    initial_temperature = float(sa.get("initial_temperature", 0))
    cooling_rate = float(sa.get("cooling_rate", 1))
    minimum_temperature = float(sa.get("minimum_temperature", 0))
    final_temperature = max(
        minimum_temperature,
        initial_temperature * cooling_rate**sa_iterations,
    )
    population = int(ga.get("population_size", 0))
    generations = int(ga.get("generations", 0))
    crossover_rate = float(ga.get("crossover_rate", 0))
    mutation_rate = float(ga.get("mutation_rate", 0))
    elite_count = ga.get("elite_count", "?")
    tournament_size = ga.get("tournament_size", "?")
    return [
        (
            f"Recocido simulado: {sa_iterations} pasos, temperatura inicial "
            f"{initial_temperature:g}, enfriamiento {cooling_rate:g}, temperatura mínima "
            f"{minimum_temperature:g} y temperatura calculada al final del presupuesto "
            f"≈ {final_temperature:.3g}. El máximo de pasos limita el tiempo; más pasos dan más "
            "oportunidades de explorar, pero aumentan el tiempo. Una temperatura inicial alta "
            "acepta más movimientos peores al principio; el enfriamiento gradual reduce esa "
            "probabilidad para estabilizar la búsqueda. La temperatura mínima evita que el "
            "cálculo de aceptación se enfríe sin límite. La temperatura inicial se mantiene "
            "común para comparar, aunque las escalas de objetivo cambien entre problemas; por "
            "eso no se afirma que sea óptima para todos."
        ),
        (
            f"Algoritmo genético: población {population}, {generations} generaciones, "
            f"cruza {crossover_rate:g}, mutación {mutation_rate:g}, elitismo {elite_count} "
            f"y torneo de {tournament_size}. La población y generaciones "
            f"definen un presupuesto básico de {population * (generations + 1)} evaluaciones "
            "por etapa genética; híbridos y búsqueda de menos colores pueden ejecutar varias "
            "etapas. Estos presupuestos moderados equilibran diversidad y costo. "
            f"Cruza {crossover_rate:g} prioriza recombinación; mutación {mutation_rate:g} añade cambios ocasionales y "
            f"elitismo {elite_count} protege los mejores individuos (el valor por defecto es 2 "
            "y se limita en poblaciones pequeñas para dejar lugar a nuevos descendientes). "
            f"El torneo de {tournament_size} da selección sin "
            "hacerla totalmente determinista. El efecto real depende del largo del cromosoma."
        ),
        (
            f"Backtracking: máximo {parameters.get('backtracking_node_limit', '?')} nodos "
            "explorados. El tope hace controlable el tiempo en instancias combinatorias grandes. "
            "Es un límite de búsqueda, no una meta de calidad: si se "
            "alcanza, el método informa truncamiento y no certifica optimalidad. Si termina "
            "sin truncarse y el problema requiere optimización, la búsqueda y sus cotas pueden "
            "certificar el óptimo."
        ),
    ]


def run_experiments(
    problems: list[str],
    runs: int,
    base_seed: int,
    sa: SAConfig,
    ga: GAConfig,
    node_limit: int,
    knapsack: KnapsackInstance | list[KnapsackInstance],
    graphs: list[GraphInstance],
) -> list[dict[str, Any]]:
    if runs < 1:
        raise ValueError("runs must be at least 1.")
    records: list[dict[str, Any]] = []
    instances: list[tuple[str, str, Any, int]] = []
    if "queens" in problems:
        instances.extend(("queens", f"N={size}", size, size) for size in (8, 50, 100))
    if "knapsack" in problems:
        knapsack_instances = [knapsack] if isinstance(knapsack, KnapsackInstance) else knapsack
        instances.extend(
            ("knapsack", item.instance_id, item, len(item.weights))
            for item in knapsack_instances
        )
    if "coloring" in problems:
        for graph in graphs:
            instances.append(
                ("coloring", f"n={graph.node_count}_seed={graph.seed}", graph, graph.node_count)
            )
    if "eggholder" in problems:
        instances.append(("eggholder", "[-512,512]^2", None, 2))

    completed = 0
    total_runs = len(instances) * runs
    for instance_index, (problem, instance_name, instance, size) in enumerate(instances):
        for run_index in range(1, runs + 1):
            seed_base = base_seed + instance_index * 100_003 + run_index * 1_009
            methods: list[tuple[str, int, SolverResult, float]] = []
            if problem == "queens":
                sa_seed, ga_seed, hybrid_seed = seed_base + 11, seed_base + 23, seed_base + 37
                methods.append(
                    (
                        "Backtracking",
                        seed_base,
                        *_timed(lambda: solve_queens_backtracking(size, node_limit)),
                    )
                )
                sa_result, sa_time = _timed(
                    lambda: queens_sa(size, sa, random.Random(sa_seed))
                )
                methods.append(("Simulated Annealing", sa_seed, sa_result, sa_time))
                ga_result, ga_time = _timed(
                    lambda: queens_ga(size, ga, random.Random(ga_seed))
                )
                methods.append(("Genetic Algorithm", ga_seed, ga_result, ga_time))

                def queens_neighbor(state: tuple[int, ...], rng: random.Random) -> tuple[int, ...]:
                    candidate = list(state)
                    left, right = rng.sample(range(size), 2)
                    candidate[left], candidate[right] = candidate[right], candidate[left]
                    return tuple(candidate)

                hybrid_result, hybrid_time = _timed(
                    lambda: hybrid_ga_sa(
                        ga_result,
                        queen_conflicts,
                        queens_neighbor,
                        sa,
                        random.Random(hybrid_seed),
                    )
                )
                hybrid_result.evaluations += ga_result.evaluations
                hybrid_result.iterations += ga_result.iterations
                hybrid_result.convergence = ga_result.convergence + hybrid_result.convergence[1:]
                hybrid_result.feasible = hybrid_result.objective == 0
                methods.append(("Hybrid GA+SA", hybrid_seed, hybrid_result, hybrid_time + ga_time))

            elif problem == "knapsack":
                methods.append(
                    (
                        "Backtracking",
                        seed_base,
                        *_timed(lambda: solve_knapsack_backtracking(instance, node_limit)),
                    )
                )
                sa_seed, ga_seed, hybrid_seed = seed_base + 11, seed_base + 23, seed_base + 37
                sa_result, sa_time = _timed(
                    lambda: knapsack_sa(instance, sa, random.Random(sa_seed))
                )
                methods.append(("Simulated Annealing", sa_seed, sa_result, sa_time))
                ga_result, ga_time = _timed(
                    lambda: knapsack_ga(instance, ga, random.Random(ga_seed))
                )
                methods.append(("Genetic Algorithm", ga_seed, ga_result, ga_time))

                def knapsack_neighbor(
                    state: tuple[int, ...], rng: random.Random
                ) -> tuple[int, ...]:
                    candidate = list(state)
                    index = rng.randrange(len(candidate))
                    candidate[index] = 1 - candidate[index]
                    return _repair_knapsack(candidate, instance)

                hybrid_result, hybrid_time = _timed(
                    lambda: hybrid_ga_sa(
                        ga_result,
                        lambda state: -float(knapsack_totals(state, instance)[1]),
                        knapsack_neighbor,
                        sa,
                        random.Random(hybrid_seed),
                    )
                )
                hybrid_result.evaluations += ga_result.evaluations
                hybrid_result.iterations += ga_result.iterations
                hybrid_result.convergence = ga_result.convergence + hybrid_result.convergence[1:]
                hybrid_result.feasible = knapsack_totals(hybrid_result.solution, instance)[0] <= instance.capacity
                methods.append(("Hybrid GA+SA", hybrid_seed, hybrid_result, hybrid_time + ga_time))

            elif problem == "coloring":
                graph: GraphInstance = instance
                methods.append(
                    (
                        "Backtracking",
                        seed_base,
                        *_timed(lambda: solve_coloring_min_backtracking(graph, node_limit)),
                    )
                )
                sa_seed, ga_seed, hybrid_seed = seed_base + 11, seed_base + 23, seed_base + 37
                sa_result, sa_time = _timed(
                    lambda: coloring_sa_min(graph, sa, random.Random(sa_seed))
                )
                methods.append(("Simulated Annealing", sa_seed, sa_result, sa_time))
                ga_result, ga_time = _timed(
                    lambda: coloring_ga_min(graph, ga, random.Random(ga_seed))
                )
                methods.append(("Genetic Algorithm", ga_seed, ga_result, ga_time))
                hybrid_result, hybrid_time = _timed(
                    lambda: coloring_hybrid_min(
                        graph,
                        sa,
                        ga,
                        random.Random(hybrid_seed),
                    )
                )
                methods.append(("Hybrid GA+SA", hybrid_seed, hybrid_result, hybrid_time))

            else:
                sa_seed, ga_seed, hybrid_seed = seed_base + 11, seed_base + 23, seed_base + 37
                sa_result, sa_time = _timed(
                    lambda: eggholder_sa(sa, random.Random(sa_seed))
                )
                methods.append(("Simulated Annealing", sa_seed, sa_result, sa_time))
                ga_result, ga_time = _timed(
                    lambda: eggholder_ga(ga, random.Random(ga_seed))
                )
                methods.append(("Genetic Algorithm", ga_seed, ga_result, ga_time))

                def egg_neighbor(
                    state: tuple[float, float], rng: random.Random
                ) -> tuple[float, float]:
                    return tuple(
                        max(-512.0, min(512.0, value + rng.gauss(0, 10)))
                        for value in state
                    )  # type: ignore[return-value]

                hybrid_result, hybrid_time = _timed(
                    lambda: hybrid_ga_sa(
                        ga_result,
                        eggholder,
                        egg_neighbor,
                        sa,
                        random.Random(hybrid_seed),
                    )
                )
                hybrid_result.evaluations += ga_result.evaluations
                hybrid_result.iterations += ga_result.iterations
                hybrid_result.convergence = ga_result.convergence + hybrid_result.convergence[1:]
                methods.append(("Hybrid GA+SA", hybrid_seed, hybrid_result, hybrid_time + ga_time))

            for method_data in methods:
                algorithm, method_seed, result, duration = _unpack_method(method_data)
                weight = value = conflicts = None
                quality = result.objective
                source = "built-in problem definition"
                instance_data: dict[str, Any]
                if problem == "knapsack":
                    weight, value = knapsack_totals(result.solution, instance)
                    quality = float(value)
                    source = instance.source
                    instance_data = {
                        "capacity": instance.capacity,
                        "weights": instance.weights,
                        "values": instance.values,
                    }
                elif problem in ("queens", "coloring"):
                    if problem == "queens":
                        conflicts = int(result.objective) if result.objective != float("inf") else None
                        instance_data = {"n": size}
                    else:
                        conflicts = result.conflicts
                        source = instance.source
                        instance_data = {
                            "nodes": instance.node_count,
                            "edges": sum(map(len, instance.adjacency)) // 2,
                            "generation_seed": instance.seed,
                            "colors_used": result.objective if result.feasible else None,
                            "sha256": _graph_digest(instance),
                        }
                else:
                    instance_data = {"bounds": [-512, 512], "dimensions": 2}
                records.append(
                    {
                        "problem": problem,
                        "instance": instance_name,
                        "instance_source": source,
                        "instance_data": json.dumps(instance_data),
                        "instance_size": size,
                        "algorithm": algorithm,
                        "run": run_index,
                        "seed": method_seed,
                        "objective": result.objective,
                        "quality": quality,
                        "feasibility_status": (
                            "Factible: respeta las restricciones"
                            if result.feasible
                            else "No factible: incumple al menos una restricción"
                        ),
                        "quality_gap_percent": "",
                        "quality_status": "Pendiente de comparación con las demás soluciones factibles",
                        "evaluations": result.evaluations,
                        "iterations": result.iterations,
                        "time_seconds": round(duration, 6),
                        "feasible": result.feasible,
                        "limit_reached": result.limit_reached,
                        "conflicts": conflicts,
                        "total_weight": weight,
                        "total_value": value,
                        "solution": json.dumps(result.solution, ensure_ascii=False),
                        "parameters": json.dumps(
                            {
                                **result_parameters(sa, ga, node_limit),
                                "graph_objective": "minimum colors among valid colorings"
                                if problem == "coloring"
                                else None,
                            },
                            sort_keys=True,
                        ),
                        "_convergence": result.convergence,
                    }
                )
            completed += 1
            progress = completed / total_runs
            print(
                f"Avance {completed}/{total_runs} ({progress:.0%}) | "
                f"{PROBLEM_LABELS.get(problem, problem)} | {_friendly_instance(problem, instance_name)} | "
                f"repetición {run_index}/{runs} lista",
                flush=True,
            )
    _annotate_record_quality(records)
    return records


def save_outputs(records: list[dict[str, Any]], output_dir: str | Path) -> None:
    target = Path(output_dir)
    target.mkdir(parents=True, exist_ok=True)
    summary = summarize(records)
    _annotate_record_quality(records)
    _save_csv(records, target / "results.csv")
    _save_csv(summary, target / "summary.csv")
    _save_excel(records, summary, target / "results.xlsx")
    _save_convergence(records, target)
    _save_report(records, summary, target / "report.md")
    parameters = json.loads(records[0]["parameters"]) if records else {}
    _save_report_pdf(summary, parameters, target / "report.pdf")
    _save_dashboard(summary, parameters, target / "report.html")


def summarize(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        grouped[(record["problem"], record["instance"], record["algorithm"])].append(record)
    summaries: list[dict[str, Any]] = []
    for (problem, instance, algorithm), group in sorted(grouped.items()):
        objectives = [float(record["objective"]) for record in group if record["objective"] is not None]
        times = [float(record["time_seconds"]) for record in group]
        evaluations = [int(record["evaluations"]) for record in group]
        qualities = [float(record["quality"]) for record in group if record["quality"] is not None]
        feasible_qualities = [
            float(record["quality"])
            for record in group
            if record["feasible"] and record["quality"] is not None
        ]
        best_quality = (max if problem == "knapsack" else min)(qualities) if qualities else ""
        best_feasible_quality = (
            (max if problem == "knapsack" else min)(feasible_qualities)
            if feasible_qualities
            else ""
        )
        summaries.append(
            {
                "problem": problem,
                "instance": instance,
                "algorithm": algorithm,
                "runs": len(group),
                "mean_objective": statistics.fmean(objectives) if objectives else "",
                "std_objective": statistics.stdev(objectives) if len(objectives) > 1 else 0.0,
                "best_objective": min(objectives) if objectives else "",
                "mean_quality": statistics.fmean(qualities) if qualities else "",
                "best_quality": best_quality,
                "mean_feasible_quality": (
                    statistics.fmean(feasible_qualities) if feasible_qualities else ""
                ),
                "best_feasible_quality": best_feasible_quality,
                "mean_time_seconds": statistics.fmean(times),
                "std_time_seconds": statistics.stdev(times) if len(times) > 1 else 0.0,
                "mean_evaluations": statistics.fmean(evaluations),
                "feasible_rate": sum(bool(record["feasible"]) for record in group) / len(group),
                "truncated_runs": sum(bool(record["limit_reached"]) for record in group),
            }
        )
    for problem, _instance, rows in _summary_groups(summaries):
        eligible = [
            row for row in rows
            if row["mean_feasible_quality"] != "" and row["feasible_rate"] > 0
        ]
        if not eligible:
            for row in rows:
                row["quality_gap_percent"] = ""
                row["quality_status"] = "Sin solución factible; calidad no comparable"
            continue
        reference = (
            max if problem == "knapsack" else min
        )(float(row["mean_feasible_quality"]) for row in eligible)
        for row in rows:
            if row not in eligible:
                row["quality_gap_percent"] = ""
                row["quality_status"] = "Sin solución factible; calidad no comparable"
                continue
            gap = _quality_gap(problem, float(row["mean_feasible_quality"]), reference)
            row["quality_gap_percent"] = round(gap, 6)
            row["quality_status"] = _quality_status(problem, gap, aggregate=True)
    return summaries


def _friendly_instance(problem: str, instance: str) -> str:
    if problem == "queens":
        return f"tablero de {instance.removeprefix('N=')} reinas"
    if problem == "knapsack":
        return f"instancia {instance.replace('kaggle_opcion1_fila_', 'fila ')}"
    if problem == "coloring":
        node_count, seed = instance.split("_seed=", maxsplit=1)
        return f"grafo de {node_count.removeprefix('n=')} nodos, semilla {seed}"
    if problem == "eggholder":
        return "función de prueba en dos dimensiones"
    return instance


def _summary_groups(
    summaries: list[dict[str, Any]],
) -> list[tuple[str, str, list[dict[str, Any]]]]:
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in summaries:
        grouped[(row["problem"], row["instance"])].append(row)
    return [
        (problem, instance, sorted(rows, key=lambda row: row["algorithm"]))
        for (problem, instance), rows in sorted(grouped.items())
    ]


def _insight_lines(summaries: list[dict[str, Any]]) -> list[str]:
    lines: list[str] = []
    for problem, instance, rows in _summary_groups(summaries):
        eligible = [
            row
            for row in rows
            if row["mean_feasible_quality"] != "" and row["feasible_rate"] > 0
        ]
        label = f"{PROBLEM_LABELS.get(problem, problem)} — {_friendly_instance(problem, instance)}"
        if not eligible:
            lines.append(f"- **{label}:** ningún método obtuvo una solución factible.")
            continue

        best_value = (
            max if problem == "knapsack" else min
        )(float(row["mean_feasible_quality"]) for row in eligible)
        best_methods = [
            row["algorithm"]
            for row in eligible
            if abs(float(row["mean_feasible_quality"]) - best_value) < 1e-9
        ]
        best_methods_display = [ALGORITHM_LABELS.get(name, name) for name in best_methods]
        fastest = min(eligible, key=lambda row: float(row["mean_time_seconds"]))
        fastest_display = ALGORITHM_LABELS.get(fastest["algorithm"], fastest["algorithm"])
        best_name = (
            f"empate entre {', '.join(best_methods_display)}"
            if len(best_methods_display) > 1
            else best_methods_display[0]
        )
        sentence = (
            f"- **{label}:** mejor promedio factible: {best_name} "
            f"({_format_quality(problem, best_value)}; objetivo: {_goal_label(problem)}). "
            f"Más rápido entre los métodos factibles: {fastest_display} "
            f"({_format_seconds(float(fastest['mean_time_seconds']))} de promedio)."
        )
        if any(row["truncated_runs"] for row in rows):
            sentence += " Backtracking alcanzó su límite en al menos una corrida; eso no prueba optimalidad."
        if any(float(row["feasible_rate"]) < 1 for row in rows):
            sentence += " Algún método no encontró solución factible en todas sus corridas."
        lines.append(sentence)
    return lines


def print_results_summary(summaries: list[dict[str, Any]]) -> None:
    print("\nRESUMEN EN LENGUAJE SENCILLO")
    print(
        "Factible = respeta las reglas; calidad = comparación solo entre las soluciones "
        "factibles. La brecha usa la mejor calidad observada, no implica optimalidad."
    )
    for line in _insight_lines(summaries):
        print(line.removeprefix("- ").replace("**", ""))
    print("\nPOR QUÉ CAMBIA EL COMPORTAMIENTO SEGÚN EL PROBLEMA")
    for title, _representation, behavior in PROBLEM_EXPLANATIONS:
        first_sentence = behavior.split(". ", maxsplit=1)[0].rstrip(".")
        print(f"- {title}: {first_sentence}.")
    print("La justificación detallada de cada representación y parámetro está en report.md, report.pdf y report.html.")


def _timed(operation: Callable[[], SolverResult]) -> tuple[SolverResult, float]:
    start = time.perf_counter()
    result = operation()
    return result, time.perf_counter() - start


def _graph_digest(graph: GraphInstance) -> str:
    digest = hashlib.sha256()
    for node, neighbors in enumerate(graph.adjacency):
        for neighbor in sorted(neighbors):
            if node < neighbor:
                digest.update(f"{node},{neighbor}\n".encode("ascii"))
    return digest.hexdigest()


def _repair_knapsack(bits: list[int], instance: KnapsackInstance) -> tuple[int, ...]:
    selected = list(bits)
    total_weight = sum(bit * weight for bit, weight in zip(selected, instance.weights))
    while total_weight > instance.capacity:
        candidates = [index for index, bit in enumerate(selected) if bit]
        index = min(candidates, key=lambda i: value_density(instance.values[i], instance.weights[i]))
        selected[index] = 0
        total_weight -= instance.weights[index]
    return tuple(selected)


def _unpack_method(data: tuple[Any, ...]) -> tuple[str, int, SolverResult, float]:
    if len(data) == 3:
        name, seed, timed_result = data
        result, duration = timed_result
        return name, seed, result, duration
    name, seed, result, duration = data
    return name, seed, result, duration


def _save_csv(rows: list[dict[str, Any]], path: Path) -> None:
    clean_rows = [{key: value for key, value in row.items() if not key.startswith("_")} for row in rows]
    if not clean_rows:
        return
    fieldnames = list(clean_rows[0])
    if path.name == "results.csv":
        fieldnames = RESULT_FIELDS
        clean_rows = [{key: row.get(key) for key in fieldnames} for row in clean_rows]
    else:
        fieldnames = list(clean_rows[0])
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(clean_rows)


def _save_excel(records: list[dict[str, Any]], summary: list[dict[str, Any]], path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as error:
        raise RuntimeError("Excel export requires openpyxl; install dependencies from requirements.txt.") from error

    workbook = Workbook()
    for title, rows in (("Results", records), ("Summary", summary)):
        sheet = workbook.active if title == "Results" else workbook.create_sheet(title)
        sheet.title = title
        clean_rows = [
            {key: value for key, value in row.items() if not key.startswith("_")}
            for row in rows
        ]
        if not clean_rows:
            continue
        headers = list(clean_rows[0])
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="234F7E")
        for row in clean_rows:
            sheet.append([row.get(header) for header in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            max_length = min(max(len(str(cell.value or "")) for cell in column) + 2, 48)
            sheet.column_dimensions[column[0].column_letter].width = max_length
    workbook.save(path)


def _save_convergence(records: list[dict[str, Any]], output_dir: Path) -> None:
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError as error:
        raise RuntimeError("Plot export requires matplotlib; install dependencies from requirements.txt.") from error

    instances = sorted({(record["problem"], record["instance"]) for record in records})
    for problem, instance in instances:
        grouped: dict[str, list[list[float]]] = defaultdict(list)
        for record in records:
            if (
                record["problem"] == problem
                and record["instance"] == instance
                and record.get("_convergence")
            ):
                grouped[record["algorithm"]].append(record["_convergence"])
        if not grouped:
            continue
        figure, axis = plt.subplots(figsize=(9, 5))
        for algorithm, curves in sorted(grouped.items()):
            length = min(map(len, curves))
            if not length:
                continue
            mean_curve = [
                statistics.fmean(float(curve[index]) for curve in curves)
                for index in range(length)
            ]
            axis.plot(range(length), mean_curve, label=algorithm)
        axis.set(
            title=f"Convergencia: {problem} — {instance}",
            xlabel="Iteración / generación",
            ylabel="Mejor objetivo",
        )
        axis.grid(True, alpha=0.25)
        axis.legend()
        figure.tight_layout()
        safe_instance = re.sub(r"[^a-zA-Z0-9_-]+", "_", instance).strip("_")
        figure.savefig(output_dir / f"convergence_{problem}_{safe_instance}.png", dpi=150)
        plt.close(figure)


def _save_report(records: list[dict[str, Any]], summaries: list[dict[str, Any]], path: Path) -> None:
    parameters = json.loads(records[0]["parameters"]) if records else {}
    has_synthetic_instances = any(
        "synthetic" in str(record["instance_source"]).lower() for record in records
    )
    lines = [
        "# Resultados del experimento",
        "",
        "## Resumen para empezar",
        "",
        "Este reporte compara cuatro formas de buscar soluciones. No necesitas conocer los algoritmos para leer el resumen: cada problema indica qué significa un resultado bueno, cuánto tardó cada método y si encontró una solución válida.",
        f"- Se completaron **{len(records)} ejecuciones** para **{len(summaries)} combinaciones** de problema, instancia y método.",
        "- «Mejor promedio factible» compara únicamente ejecuciones que encontraron una solución válida; «más rápido» compara el tiempo promedio de métodos que sí encontraron alguna solución válida.",
        "",
        "### ¿Qué método destacó?",
        "",
        *_insight_lines(summaries),
        "",
        "## Primero: factibilidad y calidad no son lo mismo",
        "",
        "- **Factibilidad** responde «¿respeta las reglas?». Mochila: peso total no supera la capacidad. Reinas: cero ataques. Coloreado: cero aristas en conflicto. Eggholder: las dos coordenadas están dentro del dominio permitido.",
        "- **Calidad** responde «entre las soluciones factibles, ¿qué tan buena es?». Se mide como valor total en mochila (más es mejor), número de colores en grafos (menos es mejor), y valor de la función Eggholder (menos es mejor). Para reinas, toda solución factible tiene exactamente cero conflictos, por lo que todas las factibles empatan en calidad.",
        "- **Mejor calidad factible observada** es una referencia empírica: la mejor media de las corridas válidas de los métodos probados en esta instancia. La desventaja porcentual compara cada método con esa referencia; no es un porcentaje de optimalidad salvo que exista certificado exacto.",
        "- **Importante:** una solución puede ser factible y de baja calidad (por ejemplo, una mochila válida que deja mucho valor fuera). Una solución candidata de alta calidad pero no factible no cuenta como solución. Una búsqueda truncada no certifica que su mejor resultado sea óptimo.",
        "",
        "## Por qué se eligieron estas representaciones",
        "",
        *[
            f"### {title}\n\n**Representación y motivo:** {representation}\n\n**Qué comportamiento puede producir:** {behavior}"
            for title, representation, behavior in PROBLEM_EXPLANATIONS
        ],
        "",
        "## Por qué estos parámetros",
        "",
        "Los valores se comparten entre instancias para que la comparación sea reproducible y sencilla de ejecutar; no se presentan como parámetros óptimos universales. Un mismo presupuesto puede alcanzar para una instancia pequeña y resultar insuficiente para otra grande o más irregular.",
        "",
        *[f"- {explanation}" for explanation in _parameter_explanations(parameters)],
        "",
        "Configuración exacta usada (también se conserva en cada fila de los datos):",
        "",
        f"```json\n{parameters}\n```",
        "",
        "## Cómo interpretar el comportamiento por problema",
        "",
        "Las explicaciones anteriores describen mecanismos del código, no garantizan un resultado. Para explicar una corrida concreta, contrasta la factibilidad, la calidad relativa, las evaluaciones, el tiempo, las curvas y los truncamientos de la tabla siguiente. Las conclusiones de «¿Qué método destacó?» son observaciones de esta ejecución, no reglas universales.",
        "",
        "## Datos y limitaciones",
        "",
        (
            "- Se usaron instancias sintéticas reproducibles; sustitúyalas por las instancias oficiales, en particular por la mochila de la tarea anterior."
            if has_synthetic_instances
            else "- Se usaron las filas de mochila y los generadores de grafos especificados en el cuaderno; `results.csv` registra el origen, las semillas y los datos de cada instancia."
        ),
        "- Las filas de `results.csv` incluyen `instance_source` e `instance_data`; el SHA-256 permite comprobar la identidad de un grafo.",
        "- Una ejecución truncada de backtracking no demuestra ausencia de solución ni optimalidad.",
        "- La desviación estándar resume la variabilidad entre ejecuciones; no equivale a una prueba de significancia.",
        "",
        "## Resultados por método",
        "",
        "La columna «Factibilidad» es la proporción de corridas que respetaron las restricciones, no una medida de calidad.",
        "",
        "| Problema | Instancia | Método | Corridas | Mejor factible | Promedio factible | Calidad vs. referencia | Factibilidad | Tiempo medio (s) | Truncadas |",
        "|---|---|---|---:|---:|---:|---|---:|---:|---:|",
    ]
    for row in summaries:
        lines.append(
            f"| {PROBLEM_LABELS.get(row['problem'], row['problem'])} | {row['instance']} | "
            f"{ALGORITHM_LABELS.get(row['algorithm'], row['algorithm'])} | {row['runs']} | "
            f"{_format_quality(row['problem'], row['best_feasible_quality'])} | "
            f"{_format_quality(row['problem'], row['mean_feasible_quality'])} | "
            f"{row['quality_status']} | {row['feasible_rate']:.1%} | "
            f"{row['mean_time_seconds']:.4f} | "
            f"{row['truncated_runs']} |"
        )
    lines.extend(
        [
            "",
            "## Visualizaciones",
            "",
            "Se genera una curva independiente por instancia para evitar mezclar escalas de tamaños diferentes. Las curvas sirven para interpretar convergencia, no como prueba estadística.",
            "",
            "## Guía de interpretación",
            "",
            "- «Mejor resultado factible» es el resultado más destacado de las corridas válidas. «Promedio factible» ayuda a ver qué tan consistente fue el método.",
            "- La mochila se compara por valor total (más alto es mejor), aunque `objective` en los archivos de datos esté guardado con signo negativo para los cálculos internos.",
            "- Una corrida truncada no demuestra que no exista una solución mejor ni que se haya encontrado el óptimo.",
            "- Compara calidad con tiempo y número de ejecuciones válidas. Con pocas repeticiones, los resultados son orientativos y no una prueba estadística.",
            "- Las curvas muestran cómo mejoró el resultado durante la búsqueda; sirven para comparar tendencias, no para demostrar que un método siempre será mejor.",
            "- Abre `report.html` para filtrar los resultados por problema, método o instancia y consultar las curvas de cada caso.",
            "",
            "## Conclusiones",
            "",
            "Las conclusiones describen únicamente las instancias y parámetros de esta ejecución. Antes de generalizarlas, considera el origen de los datos, el número de repeticiones y los límites de búsqueda.",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def _save_report_pdf(
    summaries: list[dict[str, Any]],
    parameters: dict[str, Any],
    path: Path,
) -> None:
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.backends.backend_pdf import PdfPages
    except ImportError as error:
        raise RuntimeError("PDF export requires matplotlib; install dependencies from requirements.txt.") from error

    with PdfPages(path) as pdf:
        figure = plt.figure(figsize=(11.7, 8.3))
        figure.text(0.06, 0.94, "Resultados: búsqueda y optimización", fontsize=18, weight="bold")
        cover_lines = [
            "Este informe compara cuatro métodos que buscan soluciones de forma diferente.",
            "N-Reinas: menos conflictos es mejor (cero resuelve el tablero).",
            "Mochila: mayor valor total sin exceder la capacidad es mejor.",
            "Coloreado: menos colores, siempre que no haya vecinos con el mismo color.",
            "Eggholder: menor valor de la función matemática es mejor.",
            "",
            f"Combinaciones problema/instancia/método: {len(summaries)}.",
            "«Promedio factible» resume solo las corridas con soluciones válidas.",
            "Compara ese promedio con el tiempo y la cantidad de corridas válidas.",
            "Un backtracking truncado no demuestra que la solución obtenida sea óptima.",
            "El reporte.html permite filtrar métodos y consultar gráficas por instancia.",
        ]
        for index, line in enumerate(cover_lines):
            figure.text(0.07, 0.84 - index * 0.045, line, fontsize=11)
        pdf.savefig(figure, bbox_inches="tight")
        plt.close(figure)

        explanation_pages = [
            (
                "Factibilidad y calidad: dos preguntas distintas",
                [
                    "Factibilidad pregunta si se respetan todas las restricciones. Calidad pregunta qué tan buena es la solución entre las factibles. No son sinónimos: por ejemplo, una mochila puede estar bajo la capacidad y aun así llevar poco valor.",
                    "La tabla informa por separado la tasa de corridas factibles y la brecha frente a la mejor calidad factible observada en esta instancia. Esa referencia compara métodos ejecutados; no equivale a un óptimo probado.",
                    "Para N-Reinas, una solución factible tiene cero ataques y por eso ya alcanza la calidad posible. Para mochila, mayor valor es mejor; para coloreado, menos colores es mejor una vez que no hay conflictos; para Eggholder, menor valor de función es mejor dentro del dominio.",
                    "Si un método no produjo soluciones factibles, su calidad se marca como no comparable. Si backtracking se truncó, su resultado no certifica optimalidad.",
                ],
            ),
            (
                "Presupuesto y justificación de parámetros",
                _parameter_explanations(parameters)
                + [
                    "La configuración es común a las instancias para que el experimento se pueda repetir y comparar. No se calibró independientemente para maximizar cada método en cada instancia: el tamaño y la estructura del problema cambian el efecto del mismo presupuesto.",
                    "Al explicar un resultado, considera conjuntamente calidad, factibilidad, tiempo, evaluaciones y truncamiento. Una solución más rápida pero de menor calidad, o una solución excelente hallada solo en una fracción de corridas, implica un compromiso diferente.",
                ],
            ),
        ]
        explanation_pages.extend(
            (
                f"Representación y comportamiento: {title}",
                [representation, behavior],
            )
            for title, representation, behavior in PROBLEM_EXPLANATIONS
        )
        for title, paragraphs in explanation_pages:
            _save_pdf_explanation_page(pdf, plt, title, paragraphs)

        pages = [summaries[index : index + 18] for index in range(0, len(summaries), 18)]
        for page_number, rows in enumerate(pages, 1):
            figure = plt.figure(figsize=(11.7, 8.3))
            figure.text(0.06, 0.94, "Resultados por método", fontsize=18, weight="bold")
            figure.text(
                0.06,
                0.89,
                "Factibilidad indica corridas válidas; calidad se compara con el mejor promedio factible observado.",
                fontsize=9,
            )
            columns = [
                ("Problema / instancia", 0.04),
                ("Método", 0.25),
                ("N", 0.44),
                ("Mejor", 0.49),
                ("Media", 0.58),
                ("Calidad vs mejor", 0.66),
                ("Tiempo", 0.82),
                ("Factible", 0.91),
            ]
            for title, x in columns:
                figure.text(x, 0.83, title, fontsize=8, weight="bold", ha="left")
            for index, row in enumerate(rows):
                y = 0.80 - index * 0.029
                values = [
                    f"{PROBLEM_LABELS.get(row['problem'], row['problem'])} / {row['instance']}",
                    ALGORITHM_LABELS.get(row["algorithm"], row["algorithm"]),
                    str(row["runs"]),
                    _format_quality(row["problem"], row["best_feasible_quality"]),
                    _format_quality(row["problem"], row["mean_feasible_quality"]),
                    (
                        "mejor observado"
                        if row["quality_gap_percent"] == 0
                        else (
                            f"{row['quality_gap_percent']:.1f}% brecha"
                            if row["quality_gap_percent"] != ""
                            else "no comparable"
                        )
                    ),
                    _format_seconds(float(row["mean_time_seconds"])),
                    f"{round(row['feasible_rate'] * row['runs'])}/{row['runs']}",
                ]
                for value, (_, x) in zip(values, columns):
                    figure.text(x, y, value[:25], fontsize=6.5, ha="left")
            figure.text(0.93, 0.04, f"Página {page_number}", fontsize=8, ha="right")
            pdf.savefig(figure, bbox_inches="tight")
            plt.close(figure)

        for chart_path in sorted(path.parent.glob("convergence_*.png")):
            image = plt.imread(chart_path)
            figure, axis = plt.subplots(figsize=(11.7, 8.3))
            axis.imshow(image)
            axis.axis("off")
            pdf.savefig(figure, bbox_inches="tight")
            plt.close(figure)


def _save_pdf_explanation_page(
    pdf: Any,
    plt: Any,
    title: str,
    paragraphs: list[str],
) -> None:
    lines: list[str] = []
    for paragraph in paragraphs:
        lines.extend(textwrap.wrap(paragraph, width=112))
        lines.append("")
    for page_start in range(0, len(lines), 30):
        figure = plt.figure(figsize=(11.7, 8.3))
        figure.text(0.06, 0.94, title, fontsize=16, weight="bold", wrap=True)
        page_lines = lines[page_start : page_start + 30]
        for index, line in enumerate(page_lines):
            figure.text(0.07, 0.87 - index * 0.025, line, fontsize=9, va="top")
        pdf.savefig(figure, bbox_inches="tight")
        plt.close(figure)


def _save_dashboard(
    summaries: list[dict[str, Any]],
    parameters: dict[str, Any],
    path: Path,
) -> None:
    charts = {
        (row["problem"], row["instance"]): (
            f"convergence_{row['problem']}_"
            f"{re.sub(r'[^a-zA-Z0-9_-]+', '_', row['instance']).strip('_')}.png"
        )
        for row in summaries
    }
    dashboard_rows = [
        {
            **row,
            "algorithm": ALGORITHM_LABELS.get(row["algorithm"], row["algorithm"]),
            "chart": charts[(row["problem"], row["instance"])],
        }
        for row in summaries
    ]
    explanation_markup = "\n".join(
        (
            f"<article class=\"explanation\"><h3>{html_lib.escape(title)}</h3>"
            f"<h4>Representación y motivo</h4><p>{html_lib.escape(representation)}</p>"
            f"<h4>Por qué puede comportarse así</h4><p>{html_lib.escape(behavior)}</p></article>"
        )
        for title, representation, behavior in PROBLEM_EXPLANATIONS
    )
    parameter_markup = "\n".join(
        f"<li>{html_lib.escape(line)}</li>" for line in _parameter_explanations(parameters)
    )
    explanations_section = f"""
  <section class="guide">
    <h2>Cómo leer los resultados: factibilidad y calidad</h2>
    <p><strong>Factible</strong> significa que la solución respeta las reglas del problema. <strong>Calidad</strong> mide qué tan buena es entre las soluciones factibles. No son lo mismo: una mochila puede respetar la capacidad y aun así llevar poco valor. La brecha de calidad se compara con la mejor media factible observada entre estos métodos para esta instancia; no es una prueba de optimalidad.</p>
    <p>Si no se encontró una solución factible, la calidad aparece como «no comparable». Para reinas, cero conflictos es la única calidad factible; mochila busca más valor; coloreado busca menos colores con cero conflictos; Eggholder busca un valor de función menor dentro de sus límites.</p>
    <h2>Por qué se eligieron estos parámetros</h2>
    <p>Los mismos presupuestos se aplican a las instancias para facilitar una comparación reproducible. Son presupuestos prácticos, no parámetros óptimos universales; una instancia más grande puede requerir más exploración.</p>
    <ul>{parameter_markup}</ul>
    <h2>Por qué cambia el comportamiento según el problema</h2>
    <p>Las explicaciones siguientes describen los mecanismos implementados y cómo pueden influir en los resultados. Contrástalas con factibilidad, calidad, tiempo, evaluaciones, repeticiones y truncamiento de esta ejecución.</p>
    <div class="explanations">{explanation_markup}</div>
  </section>
"""
    data = json.dumps(dashboard_rows, ensure_ascii=False)
    data = (
        data.replace("&", "\\u0026")
        .replace("<", "\\u003c")
        .replace(">", "\\u003e")
        .replace("\u2028", "\\u2028")
        .replace("\u2029", "\\u2029")
    )
    html = """<!doctype html>
<html lang="es">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Resultados del experimento</title>
  <style>
    :root { color-scheme: light; font-family: "Segoe UI", Arial, sans-serif; color: #183044; background: #f3f6fa; }
    body { margin: 0; }
    main { max-width: 1180px; margin: auto; padding: 28px 20px 48px; }
    h1 { margin-bottom: 8px; }
    .intro { max-width: 850px; line-height: 1.55; color: #40586b; }
    .filters, .cards { display: flex; flex-wrap: wrap; gap: 12px; margin: 20px 0; }
    label, .card { background: white; border: 1px solid #d8e1e9; border-radius: 8px; padding: 12px; }
    label { flex: 1 1 190px; font-weight: 600; }
    select, input { box-sizing: border-box; display: block; width: 100%; margin-top: 6px; padding: 9px; border: 1px solid #aabac8; border-radius: 5px; font: inherit; }
    .card { flex: 1 1 150px; }
    .card strong { display: block; font-size: 1.4rem; color: #135c91; }
    .table-wrap { overflow-x: auto; background: white; border: 1px solid #d8e1e9; border-radius: 8px; }
    table { width: 100%; border-collapse: collapse; min-width: 850px; }
    th, td { text-align: left; padding: 11px 10px; border-bottom: 1px solid #e4ebf0; vertical-align: top; }
    th { background: #eaf1f6; }
    .badge { display: inline-block; margin: 2px 3px 2px 0; padding: 3px 7px; border-radius: 12px; background: #e2f3e8; color: #205b35; font-size: .8rem; }
    .badge.fast { background: #e2efff; color: #174c7a; }
    .muted { color: #526777; font-size: .9rem; }
    .guide { margin: 26px 0; line-height: 1.55; }
    .explanations { display: grid; grid-template-columns: repeat(auto-fit, minmax(260px, 1fr)); gap: 12px; }
    .explanation { background: white; border: 1px solid #d8e1e9; border-radius: 8px; padding: 14px; }
    .explanation h3 { margin-top: 0; }
    .explanation h4 { margin-bottom: 4px; }
    .quality-status { display: block; margin-top: 4px; color: #40586b; font-size: .88rem; }
    details { margin-top: 8px; }
    details summary { cursor: pointer; color: #135c91; }
    img { display: block; width: min(720px, 90vw); height: auto; margin-top: 8px; }
    footer { margin-top: 22px; color: #526777; font-size: .9rem; }
    @media (max-width: 600px) { main { padding: 20px 12px; } }
  </style>
</head>
<body>
<main>
  <h1>Resultados del experimento</h1>
  <p class="intro">Explora y compara los métodos sin tener que leer una tabla técnica. Filtra por tipo de problema, método o instancia. La factibilidad indica si se respetaron las reglas; la calidad compara qué tan buena fue la solución entre las válidas.</p>
  <section class="filters" aria-label="Filtros de resultados">
    <label for="problem-filter">Tipo de problema<select id="problem-filter"><option value="">Todos</option></select></label>
    <label for="algorithm-filter">Método<select id="algorithm-filter"><option value="">Todos</option></select></label>
    <label for="search-filter">Buscar instancia<input id="search-filter" type="search" placeholder="Ej. N=50 o fila 2"></label>
  </section>
  <section class="cards" aria-live="polite">
    <div class="card"><strong id="count-combinations">0</strong>combinaciones visibles</div>
    <div class="card"><strong id="count-runs">0</strong>corridas analizadas</div>
    <div class="card"><strong id="count-truncated">0</strong>corridas truncadas</div>
  </section>
  <div class="table-wrap">
    <table>
      <thead><tr><th>Problema / instancia</th><th>Método</th><th>Calidad promedio factible</th><th>Tiempo promedio</th><th>Factibilidad</th><th>Notas y avance</th></tr></thead>
      <tbody id="results-body"></tbody>
    </table>
  </div>
  __EXPLANATIONS__
  <footer>Los resultados describen esta ejecución y sus parámetros; no garantizan que un método sea siempre mejor. Si una búsqueda exacta se truncó, no se pudo certificar que la solución encontrada sea óptima.</footer>
</main>
<script id="dashboard-data" type="application/json">__DATA__</script>
<script>
const rows = JSON.parse(document.getElementById("dashboard-data").textContent);
const problemNames = {queens:"N-Reinas", knapsack:"Mochila", coloring:"Coloreado de grafos", eggholder:"Eggholder"};
const goalNames = {queens:"menos conflictos", knapsack:"más valor", coloring:"menos colores", eggholder:"menor valor"};
const suffixes = {queens:" conflictos", knapsack:" de valor", coloring:" colores", eggholder:" puntos"};
const problemFilter = document.getElementById("problem-filter");
const algorithmFilter = document.getElementById("algorithm-filter");
const searchFilter = document.getElementById("search-filter");
function addOptions(select, values, labels) {
  [...new Set(values)].sort().forEach(value => {
    const option = document.createElement("option");
    option.value = value;
    option.textContent = labels[value] || value;
    select.append(option);
  });
}
addOptions(problemFilter, rows.map(row => row.problem), problemNames);
addOptions(algorithmFilter, rows.map(row => row.algorithm), {});
function metric(row) {
  if (row.mean_feasible_quality === "") return "Sin solución válida";
  const value = Number(row.mean_feasible_quality);
  const text = Number.isInteger(value) ? String(value) : value.toFixed(3).replace(/0+$/, "").replace(/[.]$/, "");
  return text + (suffixes[row.problem] || "");
}
function render() {
  const filtered = rows.filter(row =>
    (!problemFilter.value || row.problem === problemFilter.value) &&
    (!algorithmFilter.value || row.algorithm === algorithmFilter.value) &&
    row.instance.toLowerCase().includes(searchFilter.value.trim().toLowerCase())
  );
  const groups = new Map();
  rows.forEach(row => {
    const key = row.problem + "|" + row.instance;
    if (!groups.has(key)) groups.set(key, []);
    groups.get(key).push(row);
  });
  const highlights = new Map();
  groups.forEach((group, key) => {
    const valid = group.filter(row => row.mean_feasible_quality !== "" && Number(row.feasible_rate) > 0);
    if (!valid.length) return;
    const isKnapsack = valid[0].problem === "knapsack";
    const best = Math[isKnapsack ? "max" : "min"]( ...valid.map(row => Number(row.mean_feasible_quality)) );
    const fastest = Math.min(...valid.map(row => Number(row.mean_time_seconds)));
    highlights.set(key, {
      best: new Set(valid.filter(row => Math.abs(Number(row.mean_feasible_quality) - best) < 1e-9).map(row => row.algorithm)),
      fastest: new Set(valid.filter(row => Math.abs(Number(row.mean_time_seconds) - fastest) < 1e-9).map(row => row.algorithm))
    });
  });
  const body = document.getElementById("results-body");
  body.replaceChildren();
  filtered.forEach(row => {
    const tr = document.createElement("tr");
    const seconds = Number(row.mean_time_seconds);
    const duration = seconds < 0.001 ? "<0.001 s" : seconds.toFixed(3) + " s";
    const values = [
      (problemNames[row.problem] || row.problem) + " · " + row.instance,
      row.algorithm,
      duration,
      (Number(row.feasible_rate) * 100).toFixed(0) + "% (" + row.runs + " corridas)"
    ];
    const problemCell = document.createElement("td"); problemCell.textContent = values[0]; tr.append(problemCell);
    const algorithmCell = document.createElement("td"); algorithmCell.textContent = values[1]; tr.append(algorithmCell);
    const resultCell = document.createElement("td");
    const resultValue = document.createElement("strong"); resultValue.textContent = metric(row); resultCell.append(resultValue);
    const goal = document.createElement("div"); goal.className = "muted"; goal.textContent = "Objetivo: " + (goalNames[row.problem] || "mejor resultado"); resultCell.append(goal); tr.append(resultCell);
    const qualityStatus = document.createElement("span"); qualityStatus.className = "quality-status"; qualityStatus.textContent = row.quality_status; resultCell.append(qualityStatus);
    values.slice(2).forEach((value, index) => {
      const td = document.createElement("td");
      td.textContent = index === 1
        ? Math.round(Number(row.feasible_rate) * Number(row.runs)) + "/" + row.runs + " corridas (" + (Number(row.feasible_rate) * 100).toFixed(0) + "%)"
        : value;
      tr.append(td);
    });
    const notes = document.createElement("td");
    const key = row.problem + "|" + row.instance;
    const highlight = highlights.get(key);
    if (highlight && highlight.best.has(row.algorithm)) {
      const badge = document.createElement("span"); badge.className = "badge"; badge.textContent = "Mejor calidad promedio observada"; notes.append(badge);
    }
    if (highlight && highlight.fastest.has(row.algorithm)) {
      const badge = document.createElement("span"); badge.className = "badge fast"; badge.textContent = "Más rápido factible"; notes.append(badge);
    }
    if (Number(row.truncated_runs)) {
      const note = document.createElement("div"); note.className = "muted"; note.textContent = "Búsqueda exacta truncada en " + row.truncated_runs + " corrida(s)."; notes.append(note);
    }
    if (Number(row.feasible_rate) < 1) {
      const note = document.createElement("div"); note.className = "muted"; note.textContent = "No encontró solución válida en todas las corridas."; notes.append(note);
    }
    const details = document.createElement("details");
    const summary = document.createElement("summary"); summary.textContent = "Ver curva de avance"; details.append(summary);
    const image = document.createElement("img"); image.src = row.chart; image.alt = "Curva de avance de " + row.algorithm + " en " + row.instance; details.append(image);
    notes.append(details); tr.append(notes); body.append(tr);
  });
  document.getElementById("count-combinations").textContent = filtered.length;
  document.getElementById("count-runs").textContent = filtered.reduce((sum, row) => sum + Number(row.runs), 0);
  document.getElementById("count-truncated").textContent = filtered.reduce((sum, row) => sum + Number(row.truncated_runs), 0);
}
[problemFilter, algorithmFilter, searchFilter].forEach(control => control.addEventListener("input", render));
render();
</script>
</body>
</html>
"""
    dashboard = html.replace("__DATA__", data).replace("__EXPLANATIONS__", explanations_section)
    path.write_text(dashboard, encoding="utf-8")


def _format_number(value: Any) -> str:
    if value == "":
        return "-"
    return f"{float(value):.3f}"
